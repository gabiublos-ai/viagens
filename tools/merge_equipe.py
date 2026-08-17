#!/usr/bin/env python3
"""Atualiza a base de equipe em assets/seed.js a partir de Ace Gaming_Gestao_Colaboradores.xlsx.

Traz só o que a gestão de viagens usa — área, cargo, nível, contrato, modelo,
cidade/UF, aeroporto base e e-mail corporativo. Dados pessoais e bancários da
planilha de origem (CPF, RG, endereço, telefone, conta, PIX) NÃO são importados.

Campos que a planilha de origem não preenche (hoje, o gestor direto) são mantidos
como já estavam.

Uso: python3 tools/merge_equipe.py caminho/Ace\\ Gaming_Gestao_Colaboradores.xlsx
"""
import io
import json
import os
import sys
import unicodedata
import warnings

warnings.filterwarnings("ignore")
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SEED = os.path.join(RAIZ, "assets", "seed.js")

# Coluna da planilha de origem → campo do app
MAPA = {
    "Área": "area",
    "Cargo": "cargo",
    "Nível": "nivel",
    "Nome Gestor": "gestor",
    "Tipo Contrato": "contrato",
    "Cidade": "cidade",
    "UF": "uf",
    "Aeroporto": "aeroportoBase",
    "E-mail corporativo": "email",
    "Modelo de trabalho": "modelo",
}


def norma(s):
    return unicodedata.normalize("NFD", str(s or "").strip().lower()).encode("ascii", "ignore").decode()


def texto(v):
    if v is None:
        return ""
    return str(v).strip()


# Siglas que continuam em caixa alta ao normalizar um cargo digitado em CAIXA ALTA.
SIGLAS = {"PM", "PO", "QA", "CX", "UX", "UI", "UX/UI", "AI", "CRM", "MKT", "COS", "COO", "CEO",
          "CTO", "CFO", "RH", "TI", "SDR", "BI", "I", "II", "III", "IV"}
MINUSCULAS = {"de", "da", "do", "das", "dos", "e", "em", "para"}


def cargo_legivel(valor):
    """'DESENVOLVEDOR FULL STACK' → 'Desenvolvedor Full Stack'. Preserva siglas.

    Só age quando a origem veio inteira em caixa alta com mais de uma palavra —
    cargos já escritos normalmente passam intactos.
    """
    if not valor or valor != valor.upper() or len(valor.split()) < 2:
        return valor
    partes = []
    for i, palavra in enumerate(valor.split()):
        if palavra in SIGLAS or palavra.strip("-/") in SIGLAS:
            partes.append(palavra)
        elif i and palavra.lower() in MINUSCULAS:
            partes.append(palavra.lower())
        else:
            partes.append(palavra.capitalize())
    return " ".join(partes)


def carrega_seed():
    bruto = io.open(SEED, encoding="utf-8").read()
    cabecalho, corpo = bruto.split("window.SEED = ", 1)
    return cabecalho, json.loads(corpo.rstrip().rstrip(";"))


def le_planilha(caminho):
    ws = openpyxl.load_workbook(caminho, data_only=True)["Colaboradores"]
    # O cabeçalho fica na linha 3; as duas primeiras são título e subtítulo.
    cabecalho = [c.value for c in ws[3]]
    indices = {nome: cabecalho.index(nome) for nome in MAPA if nome in cabecalho}
    faltando = set(MAPA) - set(indices)
    if faltando:
        print("Aviso: colunas ausentes na planilha:", ", ".join(sorted(faltando)))

    registros = []
    for linha in ws.iter_rows(min_row=4, values_only=True):
        nome = texto(linha[cabecalho.index("Nome completo")])
        if not nome:
            continue
        reg = {"nome": nome}
        for coluna, campo in MAPA.items():
            if coluna in indices:
                reg[campo] = texto(linha[indices[coluna]])
        reg["cargo"] = cargo_legivel(reg.get("cargo", ""))
        registros.append(reg)
    return registros


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1

    cabecalho, seed = carrega_seed()
    atuais = {norma(c["nome"]): c for c in seed["colaboradores"]}
    novos = le_planilha(sys.argv[1])

    incluidos, alterados = [], []

    for reg in novos:
        chave = norma(reg["nome"])
        alvo = atuais.get(chave)
        if not alvo:
            alvo = {"nome": reg["nome"], "status": "Ativo", "area": "", "cargo": "", "nivel": "",
                    "gestor": "", "contrato": "", "cidade": "", "uf": "", "aeroportoBase": "",
                    "email": "", "emailAlt": "", "modelo": ""}
            seed["colaboradores"].append(alvo)
            atuais[chave] = alvo
            incluidos.append(reg["nome"])

        # O e-mail anterior (@acegaming) vira o alternativo em vez de se perder.
        antigo_email = texto(alvo.get("email"))
        if reg.get("email") and antigo_email and antigo_email != reg["email"]:
            alvo["emailAlt"] = antigo_email

        mudancas = []
        for campo, valor in reg.items():
            if campo == "nome":
                continue
            # Campo em branco na origem não apaga o que já existe (ex.: gestor).
            if not valor:
                continue
            if texto(alvo.get(campo)) != valor:
                if alvo.get(campo):
                    mudancas.append("%s: %s → %s" % (campo, alvo[campo], valor))
                alvo[campo] = valor
        if mudancas and reg["nome"] not in incluidos:
            alterados.append((reg["nome"], mudancas))

    seed["colaboradores"].sort(key=lambda c: norma(c["nome"]))
    seed["baseEquipeVersao"] = seed.get("baseEquipeVersao", 0) + 1

    with io.open(SEED, "w", encoding="utf-8") as f:
        f.write(cabecalho)
        f.write("window.SEED = ")
        json.dump(seed, f, ensure_ascii=False, indent=1)
        f.write(";\n")

    print("Base de equipe: %d pessoas (versão %d)" % (len(seed["colaboradores"]), seed["baseEquipeVersao"]))
    print("\n%d incluídos:" % len(incluidos))
    for n in incluidos:
        print("  +", n)
    print("\n%d com dados alterados:" % len(alterados))
    for nome, mudancas in alterados:
        print("  ~ %s" % nome)
        for m in mudancas:
            print("      %s" % m)
    return 0


if __name__ == "__main__":
    sys.exit(main())
