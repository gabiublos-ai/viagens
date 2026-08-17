#!/usr/bin/env python3
"""Atualiza a base de equipe a partir de "Ace Gaming_Gestao_Headcount.xlsx".

Essa planilha é a fonte da verdade do quadro: aba **Colaboradores** para quem
está ativo, aba **Desligados** para quem saiu.

Traz só o que a gestão de viagens usa — status, área, cargo, nível, contrato,
gestor direto, modelo de trabalho, cidade/UF, aeroporto e e-mails corporativos.
Remuneração, comissão, CPF, RG, telefone, endereço, dados bancários, PIX, data
de nascimento e escolaridade NÃO são importados.

Uso:
    python3 tools/merge_headcount.py caminho/Ace\\ Gaming_Gestao_Headcount.xlsx
    python3 tools/merge_headcount.py <arquivo> --simular    # só mostra o que mudaria
"""
import io
import json
import os
import sys
import unicodedata
import warnings

warnings.filterwarnings("ignore")
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SEED = os.path.join(RAIZ, "assets", "seed.js")

# Coluna da planilha → campo do app. O "ƒ" faz parte do cabeçalho da origem.
MAPA = {
    "STATUS ƒ": "status",
    "TIPO CONTRATO ƒ": "contrato",
    "ÁREA ƒ": "area",
    "CARGO ƒ": "cargo",
    "NÍVEL ƒ": "nivel",
    "GESTOR DIRETO ƒ": "gestor",
    "MODELO DE TRABALHO": "modelo",
    "CIDADE": "cidade",
    "UF": "uf",
    "AEROPORTO": "aeroportoBase",
    "E-MAIL ACE GAMING": "email",
    "E-MAIL ANTERIOR (@apostou)": "emailAlt",
}


def norma(s):
    return unicodedata.normalize("NFD", str(s or "").strip().lower()).encode("ascii", "ignore").decode()


def texto(v):
    return "" if v is None else str(v).strip()


def acha_cabecalho(ws, obrigatoria, limite=10):
    """A planilha tem título e subtítulo antes do cabeçalho; acha a linha certa."""
    for i in range(1, limite + 1):
        valores = [texto(c.value) for c in ws[i]]
        if obrigatoria in valores:
            return i, valores
    raise SystemExit("Não achei a coluna '%s' nas %d primeiras linhas de '%s'." % (obrigatoria, limite, ws.title))


def le_colaboradores(wb):
    ws = wb["Colaboradores"]
    linha_cab, cab = acha_cabecalho(ws, "NOME COMPLETO")
    indices = {MAPA[c]: cab.index(c) for c in MAPA if c in cab}
    faltando = sorted(c for c in MAPA if c not in cab)
    if faltando:
        print("Aviso: colunas ausentes na planilha:", ", ".join(faltando))

    col_nome = cab.index("NOME COMPLETO")
    registros = []
    for linha in ws.iter_rows(min_row=linha_cab + 1, values_only=True):
        nome = texto(linha[col_nome])
        if not nome:
            continue
        reg = {"nome": nome}
        for campo, coluna in indices.items():
            reg[campo] = texto(linha[coluna])
        registros.append(reg)
    return registros


def le_desligados(wb):
    if "Desligados" not in wb.sheetnames:
        return {}
    ws = wb["Desligados"]
    try:
        linha_cab, cab = acha_cabecalho(ws, "NOME COMPLETO")
    except SystemExit:
        return {}
    col_nome = cab.index("NOME COMPLETO")
    col_dia = cab.index("DATA ÚLTIMO DIA") if "DATA ÚLTIMO DIA" in cab else None
    saidas = {}
    for linha in ws.iter_rows(min_row=linha_cab + 1, values_only=True):
        nome = texto(linha[col_nome])
        if nome:
            saidas[norma(nome)] = {"nome": nome, "ultimoDia": texto(linha[col_dia]) if col_dia is not None else ""}
    return saidas


def carrega_seed():
    bruto = io.open(SEED, encoding="utf-8").read()
    cabecalho, corpo = bruto.split("window.SEED = ", 1)
    return cabecalho, json.loads(corpo.rstrip().rstrip(";"))


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    caminho = sys.argv[1]
    simular = "--simular" in sys.argv

    wb = openpyxl.load_workbook(caminho, data_only=True)
    ativos = le_colaboradores(wb)
    desligados = le_desligados(wb)

    cabecalho, seed = carrega_seed()
    atuais = {norma(c["nome"]): c for c in seed["colaboradores"]}

    incluidos, alterados, marcados = [], [], []

    for reg in ativos:
        chave = norma(reg["nome"])
        alvo = atuais.get(chave)
        if not alvo:
            alvo = {"nome": reg["nome"], "status": "Ativo", "area": "", "cargo": "", "nivel": "",
                    "gestor": "", "contrato": "", "cidade": "", "uf": "", "aeroportoBase": "",
                    "email": "", "emailAlt": "", "modelo": ""}
            seed["colaboradores"].append(alvo)
            atuais[chave] = alvo
            incluidos.append(reg["nome"])

        mudancas = []
        for campo, valor in reg.items():
            if campo == "nome" or not valor:
                continue
            if texto(alvo.get(campo)) != valor:
                if alvo.get(campo):
                    mudancas.append("%s: %s → %s" % (campo, alvo[campo], valor))
                alvo[campo] = valor
        if mudancas and reg["nome"] not in incluidos:
            alterados.append((reg["nome"], mudancas))

    # Quem saiu fica no cadastro, marcado — as viagens passadas continuam válidas.
    for chave, saida in desligados.items():
        alvo = atuais.get(chave)
        if alvo and alvo.get("status") != "Desligado":
            alvo["status"] = "Desligado"
            marcados.append(saida["nome"] + (" (último dia %s)" % saida["ultimoDia"][:10] if saida["ultimoDia"] else ""))

    conhecidos = {norma(r["nome"]) for r in ativos} | set(desligados)
    fora = [c["nome"] for c in seed["colaboradores"] if norma(c["nome"]) not in conhecidos]

    seed["colaboradores"].sort(key=lambda c: norma(c["nome"]))
    seed["baseEquipeVersao"] = seed.get("baseEquipeVersao", 0) + 1

    print("Planilha: %d ativos · %d desligados" % (len(ativos), len(desligados)))
    print("Base do app: %d pessoas (versão %d)\n" % (len(seed["colaboradores"]), seed["baseEquipeVersao"]))

    print("%d incluídos:" % len(incluidos))
    for n in incluidos:
        print("  +", n)
    print("\n%d marcados como desligados:" % len(marcados))
    for n in marcados:
        print("  ×", n)
    print("\n%d com dados alterados:" % len(alterados))
    for nome, mudancas in alterados:
        print("  ~ %s" % nome)
        for m in mudancas:
            print("      %s" % m)
    if fora:
        print("\n%d no app e fora da planilha (mantidos, confira):" % len(fora))
        for n in fora:
            print("  ?", n)

    if simular:
        print("\n(simulação — assets/seed.js não foi alterado)")
        return 0

    with io.open(SEED, "w", encoding="utf-8") as f:
        f.write(cabecalho)
        f.write("window.SEED = ")
        json.dump(seed, f, ensure_ascii=False, indent=1)
        f.write(";\n")
    print("\nassets/seed.js atualizado.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
