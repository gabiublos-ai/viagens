#!/usr/bin/env python3
"""Extrai a base da planilha Gestao_Viagens_Corporativas para assets/seed.js."""
import json, sys, warnings, datetime
warnings.filterwarnings("ignore")
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "Gestao_Viagens_Corporativas_AceGaming_2026.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "assets/seed.js"

wb = openpyxl.load_workbook(SRC, data_only=True)

def s(v):
    if v is None: return ""
    if isinstance(v, str): return v.strip()
    return str(v).strip()

def n(v):
    if v in (None, ""): return 0
    try: return round(float(v), 2)
    except (TypeError, ValueError): return 0

def d(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return ""

# --- Colaboradores -------------------------------------------------------
ws = wb["Colaboradores"]
colaboradores = []
for r in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
    if not r[1]: continue
    colaboradores.append({
        "nome": s(r[1]), "status": s(r[2]), "area": s(r[3]), "cargo": s(r[4]),
        "nivel": s(r[5]), "gestor": s(r[6]) if s(r[6]) != "0" else "",
        "contrato": s(r[7]), "cidade": s(r[8]), "uf": s(r[9]),
        "aeroportoBase": s(r[10]), "email": s(r[11]), "emailAlt": s(r[12]),
        "modelo": s(r[13]),
    })

# --- Viagens -------------------------------------------------------------
ws = wb["Viagens"]
viagens = []
for r in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
    if not r[1]: continue
    viagens.append({
        "id": int(r[1]),
        "tipo": s(r[2]) or "Viagem",
        "colaborador": s(r[4]),
        "destino": s(r[7]),
        "aeroportoOrigem": s(r[8]),
        "dataIda": d(r[9]),
        "dataVolta": d(r[10]),
        "aereo": n(r[12]),
        "diarias": n(r[13]),
        "valorDiaria": n(r[14]),
        "alimentacao": n(r[16]),
        "transporte": n(r[17]),
        "custoAlteracao": n(r[18]),
        "status": s(r[23]) or "Fechado",
        "refId": int(r[24]) if r[24] not in (None, "") else None,
        "motivo": s(r[25]),
        "pendencias": s(r[26]),
        "obs": s(r[27]),
    })

# --- Uber (linhas cruas do relatório) ------------------------------------
ws = wb["Colar Uber"]
uber = []
for r in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
    if r[0]: uber.append(s(r[0]))

# --- Parâmetros ----------------------------------------------------------
ws = wb["Parâmetros"]
meses, categorias, statuses, servicos, tiposTransacao, kwAeroporto = [], [], [], [], [], []
deparaRaw = []
for r in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
    if s(r[1]): meses.append(s(r[1]))
    if s(r[3]): categorias.append(s(r[3]))
    if s(r[4]): statuses.append(s(r[4]))
    if s(r[5]) and s(r[5]) != "--": servicos.append(s(r[5]))
    if s(r[6]): tiposTransacao.append(s(r[6]))
    if s(r[8]) and not s(r[8]).startswith("↑"): kwAeroporto.append(s(r[8]).lower())
    if s(r[12]) and s(r[13]) and s(r[12]) != "Nome como aparece no relatório do Uber":
        deparaRaw.append({"uber": s(r[12]), "colaborador": s(r[13])})

regras = {"jantar": 70, "cafe": 30, "toleranciaUber": 1}
ws_p = wb["Parâmetros"]
regras["jantar"] = n(ws_p["L5"].value) or 70
regras["cafe"] = n(ws_p["L6"].value) or 30
regras["toleranciaUber"] = int(n(ws_p["L7"].value) or 1)

seed = {
    "versao": 1,
    "empresa": "Ace Gaming / apostou.bet.br",
    "colaboradores": colaboradores,
    "viagens": viagens,
    "uberRaw": uber,
    "parametros": {
        "meses": meses,
        "categorias": categorias,
        "statuses": statuses,
        "servicos": servicos,
        "tiposTransacao": tiposTransacao,
        "palavrasAeroporto": kwAeroporto,
        "regras": regras,
        "dePara": deparaRaw,
    },
}

with open(OUT, "w", encoding="utf-8") as f:
    f.write("/* Gerado por tools/extract_seed.py a partir da planilha original. */\n")
    f.write("window.SEED = ")
    json.dump(seed, f, ensure_ascii=False, indent=1)
    f.write(";\n")

print(f"{len(colaboradores)} colaboradores · {len(viagens)} viagens · {len(uber)} linhas Uber → {OUT}")
